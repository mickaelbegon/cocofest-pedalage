from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "docs/scientific_appendix/appendix_acceleration_ding.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tc_pr.append(shd)


def border(cell):
    tc_pr = cell._tc.get_or_add_tcPr(); borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        line = OxmlElement(f"w:{edge}"); line.set(qn("w:val"), "single")
        line.set(qn("w:sz"), "4"); line.set(qn("w:color"), "D9D9D9"); borders.append(line)
    tc_pr.append(borders)


def paragraph(doc, text="", style=None):
    result = doc.add_paragraph(style=style); result.add_run(text)
    result.paragraph_format.space_after = Pt(6); result.paragraph_format.line_spacing = 1.12
    return result


def equation(doc, text):
    result = doc.add_paragraph(); result.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = result.add_run(text); run.font.name = "Cambria Math"
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Cambria Math"); run.font.size = Pt(10.5)
    result.paragraph_format.space_after = Pt(8)


doc = Document(); section = doc.sections[0]
section.top_margin = Inches(.75); section.bottom_margin = Inches(.72)
section.left_margin = Inches(.78); section.right_margin = Inches(.78)
styles = doc.styles; styles["Normal"].font.name = "Aptos"; styles["Normal"].font.size = Pt(10)
for name, size in (("Title", 17), ("Heading 1", 13.5), ("Heading 2", 11.5)):
    styles[name].font.name = "Aptos"; styles[name].font.size = Pt(size); styles[name].font.color.rgb = RGBColor(0, 0, 0)

title = paragraph(doc, "Technical appendix Computational acceleration for FES cycling with the Ding model", "Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
properties = title._p.get_or_add_pPr(); existing = properties.find(qn("w:pBdr"))
if existing is not None: properties.remove(existing)
title_border = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom"); bottom.set(qn("w:val"), "nil")
title_border.append(bottom); properties.append(title_border)
subtitle = paragraph(doc, "Mathematical construction of reduced mechanics and local Ding-state reduction")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
paragraph(doc, "This appendix explains the two main model reductions used for FES cycling: projection of the constrained multibody model onto the crank manifold, and local reconstruction of linear Ding states. It distinguishes an exact identity of a selected discrete transcription from continuous-model equivalence and from observed solver speed. Measurement-corrected NMPC remains necessary for use with a participant.")

paragraph(doc, "1. Historical baseline and notation", "Heading 1")
paragraph(doc, "The historical reference is the FES hand-cycling formulation described by Co et al. [7]. It already includes RHO NMPC with primal warm start, IPOPT with MA57 on Linux, degree-3 Radau collocation, a fatigable Ding model, and force-length, force-velocity, and passive-force relationships. The main configuration is 30 Hz, a two-cycle window, a one-cycle shift, and five requested cycles. These are baseline elements, not later innovations.")
paragraph(doc, "For M muscles, the full state contains five Ding states per muscle, x = (Cn, F, A, Tau1, Km), and the mechanical generalized coordinates q and qdot. The physical crank coordinate is theta and omega = theta-dot. Pulse width is denoted u.")

paragraph(doc, "2. Starting point The Ding force fatigue system", "Heading 1")
paragraph(doc, "The five-state phenomenological model follows the Ding family of force and fatigue models [1-3]. Define the pulse-width recruitment and calcium sensitivity terms:")
equation(doc, "ρ(u) = 1 − exp(−(u − u₀)/uₜ),     σ(C,K) = C/(K + C)")
paragraph(doc, "Let g(theta, omega) collect the active force-length, force-velocity, and passive-force factors. The implementation used in the OCP is:")
equation(doc, "Ċ = (H(t) − C)/τc")
equation(doc, "Ḟ = g(θ,ω)[A ρ(u) σ(C,K) − F/(T + τ₂ σ(C,K))]")
equation(doc, "Ȧ = −(A − Ar)/τf + αA F,     Ṫ = −(T − Tr)/τf + αT F,     K̇ = −(K − Kr)/τf + αK F")
paragraph(doc, "The force equation remains nonlinear and coupled to mechanics. The A, Tau1, and Km equations, however, share the same recovery operator and the same forcing F. This exact structural property is what enables local state reduction; it does not make the force dynamics linear.")

paragraph(doc, "3. Reduced crank mechanics", "Heading 1")
paragraph(doc, "3.1 Contact manifold and tangent projection", "Heading 2")
paragraph(doc, "The full model has generalized coordinates q and hand-to-crank contact constraints. Offline, the physical contact problem is solved over one crank revolution, yielding a periodic contact-consistent configuration curve q = Q(theta). The reduced model therefore parameterizes the admissible contact manifold by the physical crank angle rather than fitting an arbitrary torque curve.")
equation(doc, "q = Q(θ),     q̇ = Q′(θ)ω,     q̈ = Q′(θ)ω̇ + Q″(θ)ω²")
paragraph(doc, "Starting from M(q)q-double-dot + h(q,q-dot) = tau_m + tau_ext and premultiplying by Q-prime-transpose defines the effective inertia I(theta), gravitational term G(theta), quadratic velocity term V(theta), muscle effectiveness eta_m(theta), and external-torque effectiveness eta_e(theta). The online dynamics are:")
equation(doc, "θ̇ = ω")
equation(doc, "I(θ)ω̇ = Σm ηm(θ)Fm + ηe(θ)τext − G(θ) − V(θ)ω²")
paragraph(doc, "This is a Lagrange/tangent projection of the constrained multibody equations. It removes redundant generalized coordinates while retaining the crank coordinate and speed. It is valid on the sampled contact branch and does not model departures from it.")

paragraph(doc, "3.2 Construction of the reduced numerical profile", "Heading 2")
paragraph(doc, "At sampled crank angles, the implementation evaluates the Biorbd mass matrix, nonlinear effects, muscle-length Jacobian, and Q, Q-prime, and Q-double-prime. It computes I = Q-prime-transpose M Q-prime, eta_m = minus J-length,m Q-prime, and eta_e = Q-prime,r. It also obtains normalized fibre length and fibre velocity per unit crank speed. Smooth periodic Fourier series are fitted to I, G, V, the effectiveness coefficients, and muscle geometry. Online dynamics are therefore small symbolic functions of theta, omega, and muscle forces instead of a full contact and multibody solve.")
equation(doc, "l̃m = lm(θ),     ṽm = νm(θ)ω/10")
paragraph(doc, "The original Hill force-length, force-velocity, and passive laws are evaluated from these fitted geometry profiles. The reduced profile must be regenerated and audited whenever the source biomechanical model, contact geometry, selected muscles, or Fourier order changes. Independent samples must compare projected acceleration and muscle geometry against the full constrained model.")

paragraph(doc, "3.3 Measured effect and scope", "Heading 2")
paragraph(doc, "For four muscles, the full forward model has 20 Ding states plus six mechanical states. Dynamic reduced mechanics keeps the 20 Ding states and theta, omega, reducing 26 states to 22. In the controlled 100-RHO IPOPT/Radau-3 campaign, hot median time changed from 4.595 to 1.031 s and total wall time from 615.4 to 260.9 s; the fatigue difference was 0.097 percent. These are campaign-specific results and must not be multiplied with gains from other modifications. Reduced mechanics is a model transformation, not merely a code optimization.")

paragraph(doc, "4. Local reduction of Ding states", "Heading 1")
paragraph(doc, "4.1 Continuous identities", "Heading 2")
paragraph(doc, "For z in {A, T, K}, with rest value zr and fatigue coefficient alpha-z, an integrating factor gives:")
equation(doc, "z(tk+s) = zr + exp(−s/τf)(zk − zr) + αz ∫₀ˢ exp(−(s−v)/τf)F(tk+v)dv")
paragraph(doc, "All three states share the same filtered-force convolution. If J(s) = integral from 0 to s of exp(v/tau-f) F(tk+v) dv, then z(tk+s) = zr + exp(−s/tau-f)[zk − zr + alpha-z J(s)]. This permits a continuous formulation using F and J, with A, Tau1, and Km reconstructed.")
paragraph(doc, "A more convenient local formulation keeps F and A. When alpha-A is nonzero, set beta-T = alpha-T/alpha-A, beta-K = alpha-K/alpha-A, dT = (T−Tr)−beta-T(A−Ar), and dK = (K−Kr)−beta-K(A−Ar). The force terms cancel exactly:")
equation(doc, "ḋT = −dT/τf,     ḋK = −dK/τf")
equation(doc, "T = Tr + βT(A − Ar) + dT,k exp(−s/τf),     K = Kr + βK(A − Ar) + dK,k exp(−s/τf)")
paragraph(doc, "For the periodic-node calcium input H(tk+s) = Hk exp(−s/tau-c), calcium is also reconstructible on an interval:")
equation(doc, "C(tk+s) = exp(−s/τc)[Ck + Hk s/τc]")
paragraph(doc, "A stimulation changes the forcing amplitude but not the calcium state. The reconstruction must use the real incoming RHO state, including nonzero fatigue offsets; replacing it by a periodic fixed point changes the physical initial condition.")

paragraph(doc, "4.2 Preserve the discrete transcription", "Heading 2")
paragraph(doc, "Continuous identities do not reproduce a finite-step Radau or Gauss-Legendre map exactly. For a Runge-Kutta tableau (A,b,c) and step h, preserve the OCP by eliminating only the linear equations after discretization. For five-stage Radau IIA:")
equation(doc, "C⃗ = (I + h A/τc)⁻¹[1⃗ Ck + (h/τc)A H⃗]")
equation(doc, "z⃗ = zr1⃗ + (I + h A/τf)⁻¹[1⃗(zk − zr) + h αz A F⃗]")
paragraph(doc, "After this substitution, the nonlinear muscle residual contains force-stage variables, while calcium and fatigue stages are reconstructed. This is algebraically identical to the full Radau discrete equations up to linear-solve and NLP tolerances. Radau IIA is stiffly accurate, so the last stage is the endpoint [4,5]. A five-state, five-stage muscle has 25 stage variables; this local condensation leaves five nonlinear force stages, while the physical node states may be retained for RHO transfer and constraints.")
paragraph(doc, "For Acados IRK Gauss-Legendre 4x5, eliminated C, dT, and dK profiles must be built with the Gauss-Legendre tableau, substep by substep. The endpoint is not the final stage and must be evaluated as:")
equation(doc, "vk+1 = vk + h bᵀ[r(tk + hc) − V⃗/τ]")
paragraph(doc, "Reusing Radau profiles in Gauss-Legendre would change the discrete problem.")

paragraph(doc, "4.3 Dimension, constraints, and derivatives", "Heading 2")
paragraph(doc, "With four muscles and reduced mechanics, retaining only F and A gives 10 physical states instead of 22; C, Tau1, and Km are reconstructed locally. This differs from Radau-stage condensation, which can keep all physical states at nodes while eliminating internal stage variables. The 10-state count requires fixed stimulation scheduling, calcium parameters, and incoming state in a window.")
paragraph(doc, "Every cost or constraint using a removed variable must be evaluated on its reconstruction at the same nodes or stages as in the full problem. For scaled variables, a relation y-bar = beta A-bar + delta converts bounds only after correct scaling and sign handling. The physical domain K+C > 0 and T + tau-2 sigma(C,K) > 0 remains mandatory. Reconstructed expressions must remain in the automatic-differentiation graph; treating a state-dependent profile as constant gives incorrect sensitivities. Direct-collocation formulations are known to change both robustness and computational cost [6].")

paragraph(doc, "4.4 Evidence and limitations", "Heading 2")
paragraph(doc, "On isolated, converged maps, the local ACADOS proof produced maximum reconstruction errors of 1.33e−16 for GL4x5 and 3.61e−15 for Radau IIA five-stage; selected sensitivity errors were 3.24e−12 and 9.48e−12. These establish map equivalence, not OCP performance. In the IPOPT/Radau-5 reduced RHO A/B, both variants completed 50 propagated windows; hot median time changed from 0.798 to 0.551 s (−30.9 percent). For ACADOS, the experimental local formulation completed 76 valid cycles before stopping at cycle 77; over the first 75 hot cycles the median changed from 0.455 to 0.167 s (−63.3 percent). This is promising but not a 100/100 robustness claim.")

paragraph(doc, "5. Solver comparison and benchmark protocol", "Heading 1")
paragraph(doc, "IPOPT/MA57 with five-stage Radau IIA is the fidelity reference. Acados uses multiple shooting with an IRK Gauss-Legendre 4x5 integrator. Common incoming state, seed, Ding parameters, mechanical profile, load, objective, and bounds are prerequisites to comparison. Both control sequences must be replayed with DOP853 under fixed controls. Define the continuous speed excursion:")
equation(doc, "εω = maxₜ max(0, ωmin − ω(t), ω(t) − ωmax)")
paragraph(doc, "The controlled pair had DOP853 costs of 0.01295477 for IPOPT and 0.01296194 for Acados, but both had small between-node excursions. With a common tightened fast margin, costs were 0.01383435 and 0.01384597 (Acados +0.084 percent), with excursions 0.00515 and 0.00772 rad/s. Strict ranking requires zero excursion; similar small excursions support only a qualified numerical comparison.")
paragraph(doc, "Every ablation should report build and compilation time, hot-solve time, wall time, iterations, valid-cycle fraction, common DOP853 cost, epsilon-omega, phase error, pulse-width bounds, KKT dimensions/nonzeros, and control distance. RHO-FHO tests require equal horizons and separate per-cycle and total times. Individual speedup factors must not be multiplied.")

paragraph(doc, "6. Cumulative 30 Hz ablation and frequency campaign", "Heading 1")
paragraph(doc, "The following cumulative IPOPT MA57 ablation uses 100 one-cycle RHO windows, reduced dynamic mechanics, Radau IIA degree five, and 30 uniformly spaced stimulations per crank revolution. A result is validated only if every window converges, has primal feasibility below 1e-5, and passes the mechanical audit. Cold construction, compilation, and audit time are excluded from the hot percentiles, which describe cycles 2 to 100.")

def table(doc, headers, rows, widths=None):
    value = doc.add_table(rows=1, cols=len(headers)); value.alignment = WD_TABLE_ALIGNMENT.CENTER
    value.style = "Table Grid"
    for j, header in enumerate(headers):
        cell = value.rows[0].cells[j]; cell.text = header; shade(cell, "1F4E78"); border(cell)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255); run.font.bold = True; run.font.size = Pt(8)
    for i, row in enumerate(rows):
        cells = value.add_row().cells
        for j, item in enumerate(row):
            cells[j].text = str(item); border(cells[j]); cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if i % 2: shade(cells[j], "EAF2F8")
            for p in cells[j].paragraphs:
                p.paragraph_format.space_after = Pt(1); p.paragraph_format.line_spacing = 1.0
                for run in p.runs: run.font.size = Pt(8)
        if widths:
            for j, width in enumerate(widths): cells[j].width = Inches(width)
    paragraph(doc, "")
    return value

table(doc,
    ["Case", "Incremental change", "Validated", "Median s", "P90 s", "Interpretation"],
    [
        ["K3", "Reduced mechanics", "100/100", "5.943", "6.918", "MX graph remains expensive"],
        ["K4", "Periodic-node calcium", "100/100", "5.740", "6.584", "Small pre-compilation effect"],
        ["K5", "SX graph", "100/100", "0.678", "0.798", "Largest observed acceleration"],
        ["K6", "Compiled exact Hessian", "100/100", "0.449", "0.521", "Fastest valid 30 Hz case"],
        ["K7", "Compact RHO output", "100/100", "0.465", "0.563", "Small output-retention cost"],
        ["K8", "Hard 100 us delta PW bound", "100/100", "0.920", "1.151", "Regularity constraint costs time"],
        ["K9", "Delta PW penalty, weight 0.01", "100/100", "0.789", "0.892", "Best regularized case"],
    ], [0.42, 1.2, 0.55, 0.52, 0.48, 2.85])

paragraph(doc, "SX is the largest timing change in this implementation and the compiled Hessian is beneficial after symbolic reduction. Speedup factors must not be multiplied because each row changes the symbolic graph and constrained NLP. K1 and K2 full-mechanics runs solved numerically but failed absolute crank-progress validation; their 29.739 s and 27.562 s medians are not solver comparisons. K0 degree-three Radau was intentionally excluded at 30 Hz by the independent integration gate.")

paragraph(doc, "6.1 Frequency and periodic calcium", "Heading 2")
paragraph(doc, "For a one-second cycle with Ns evenly spaced pulses, the periodic-node model uses h = 1/Ns and recomputes decay, post-pulse amplitude, and the periodic calcium fixed point at each frequency.")
equation(doc, "d = exp(−h/τc),     Cn* = d H0(h)(h/τc)/(1−d)")
paragraph(doc, "The frequency runs are therefore fixed-frequency periodic-calcium experiments, not reuses of a 30 Hz calcium trace. They represent the steady periodic regime after warm-up. A true change in frequency requires phase-resampling controls and initialising calcium at the fixed point for the new interval, or explicitly simulating the transition.")

table(doc,
    ["Frequency and case", "Solver", "Validated", "Median s", "P90 s", "Status"],
    [
        ["30 Hz K5", "MadNLP 0.10.1 MA57", "100/100", "0.925", "1.193", "Native valid comparator"],
        ["33 Hz K7", "IPOPT MA57", "100/100", "0.558", "0.681", "Valid, below one second"],
        ["35 Hz K5", "IPOPT MA57", "100/100", "0.875", "1.451", "Valid, tail above one second"],
        ["35 Hz K5", "MadNLP 0.10.1 MA57", "100/100", "1.148", "1.399", "Valid, not real-time"],
        ["40 Hz K5", "IPOPT MA57", "100/100", "1.043", "1.319", "Valid, median above one second"],
        ["40 Hz K5", "MadNLP 0.10.1 MA57", "0/100", "--", "--", "Native continuity failure"],
        ["50 Hz K10", "IPOPT MA57", "100/100", "1.488", "1.716", "Robustness point, not real-time"],
    ], [0.88, 1.3, 0.55, 0.52, 0.48, 2.3])

paragraph(doc, "The 33 Hz K7 result is the conservative high-frequency reference: it passes the mechanical audit and remains below one second at P90. The 50 Hz K10 result is robust but outside a one-second control budget. The original IPOPT MUMPS protocol did not produce a native valid 100-cycle trajectory, so no MA57 versus MUMPS speed ratio is reported. Acados remains diagnostic because of a between-node velocity excursion. Its local Ding map checks are near machine precision, but its 76-cycle stopping point does not support a 100-cycle performance claim.")

paragraph(doc, "6.2 Decision efficient validation policy", "Heading 2")
paragraph(doc, "Rather than test every solver, frequency, and graph option, use IPOPT MA57 as the certified frequency reference. Advance frequency only after feasibility passes. Test an independent solver at 30 Hz, at the last IPOPT point whose P90 is below one second, and at the first point beyond this threshold. Test MUMPS only as an IPOPT linear-solver A/B, and retain Acados only after common DOP853 replay passes. Final comparisons must report DOP853 objective, speed excursion, phase or work error, and pulse-width feasibility separately from native solver timings.")

paragraph(doc, "6.3 Terminal half-step guard at the RHO seam", "Heading 2")
paragraph(doc, "In free-cadence reduced mechanics, the terminal state of a RHO window is the fixed incoming state of the next window. The optional seam guard reuses the reduced internal cadence predictor at the terminal node, before the next window can change a control. With h one shooting interval and fomega the reduced angular acceleration, it constrains the Euler half-step prediction to the same speed interval as the internal guard.")
equation(doc, "ωhalf = ωT + (h/2) fω(θT, ωT, FT, τext)")
paragraph(doc, "The guard is available only with dynamic reduced mechanics and an active internal reduced-speed guard. It adds no decision variable and does not trigger a new NLP compilation. It is a local, differentiable Euler prediction—not a DOP853 replay and not a certificate for every constraint over the next cycle.")
table(doc,
    ["Terminal seam guard", "Certified", "Hot median / P90 s", "Validated loop wall s", "Interpretation"],
    [
        ["Off", "100/100", "0.650 / 0.898", "81.172", "Reference internal guard only"],
        ["Euler half-step on", "100/100", "0.680 / 0.871", "80.961", "Timing-neutral within this small A/B"],
    ], [0.8, 0.6, 1.0, 1.0, 2.0])
paragraph(doc, "This matched IPOPT/MA57/Radau-5 100-window reduced-dynamics A/B shows that the local seam constraint did not remove the observed one-second-scale operation. It does not prove a continuous-time or clinical safety gain. A separate diagnostic found an Euler prediction at 0.9h more conservative than the continuous minimum speed; extra Euler samples or a tightened margin are therefore not presented as certificates. The acceptance test remains a dense DOP853 replay spanning the seam and next applied cycle.")

paragraph(doc, "6.4 Two independent arms and adaptive muscle weights", "Heading 2")
paragraph(doc, "Two independent unilateral isokinetic OCPs are used to allocate a prescribed bilateral work budget, not to represent a coupled two-arm crank. Each arm has four Ding muscles, its own state history and pulse widths. A process-isolated coordinator alone enforces the work identity after certification:")
equation(doc, "WR(k) + WL(k) = 2π τbar,total")
paragraph(doc, "No state from one arm enters the other arm's compiled NLP. At cold start the manual split is the declared pair of equivalent mean torques. A capacity–fatigability alternative waits for one certified common cycle. From each maximum-PW isokinetic envelope, csm is a positive work opportunity and dsm the exact prescribed-force Ding capacity-decrement ratio. With Cs = sum(csm), qsm = csm/Cs, Ds = sum(qsm dsm), and Es = Wref/Ds, the initial right share is pR = ER/(ER + EL), clipped only by a declared minimum arm torque. Cs is an opportunity measure, not a feasibility certificate.")
paragraph(doc, "After both arms certify a block boundary, the causal work-split reference uses pR-star = rR-to-the-gain divided by (rR-to-the-gain + rL-to-the-gain), then applies smoothing, a maximum fraction step, and the minimum arm torque. It is work-conserving and updates only every K certified cycles (K=10 by default), but it does not predict the endurance optimum.")
paragraph(doc, "Each unilateral cost is sum over muscles of wm times (1 − Am/Arest,m) squared. Weight vectors have geometric mean one, relative bounds, and bounded log updates; numeric parameter binding preserves the single compiled NLP. The explicit scientific conditions are:")
table(doc,
    ["Policy", "Cold start", "Certified slow update", "Role"],
    [
        ["Unit", "wm = 1", "None", "Neutral reference"],
        ["Physio-U", "Unit; no FHO/BO", "Mechanical Shapley credit × state-conditioned Ding force challenge", "Physiological hypothesis"],
        ["Mechanical-sensitivity-squared v1", "Unit; first update waits for live envelope", "wm target proportional to bm squared", "Mechanics-only ablation"],
        ["Capacity feedback", "Unit", "Inverse local capacity-ratio feedback", "Causal reference"],
    ], [1.1, 1.45, 2.1, 1.35])
paragraph(doc, "For the mechanics-only condition, bm is normalized mechanical Shapley work credit. It intentionally does not multiply alpha-A or a fatigue decrement: Ding fatigue is already propagated in A and the objective already squares normalized fatigue. Physio-U remains distinct because it does use the state-conditioned force challenge. Missing, invalid, zero-credit, out-of-domain, or uncertified envelopes hold the incumbent weights. At due boundaries, the live envelope uses the full local Ding state (Cn, F, A, Tau1, Km), pulse-width limits, reduced isokinetic geometry, and certified terminal work; the target is projected, smoothed, and log-step capped before the numeric parameter update.")

paragraph(doc, "6.5 Isokinetic branch: prescribed kinematics before SX", "Heading 2")
paragraph(doc, "The K0–K9 table is a dynamic-mechanics lineage and remains separate: switching to a prescribed cadence changes the task, so no timing factor is propagated from it into the isokinetic branch. In that branch, prescribed kinematics is introduced immediately after reduced mechanics and before SX: K3-I0 is the reduced MX state-kinematics reference and K3-I1 eliminates the prescribed theta/omega equalities. The intended continuation is K4-I periodic calcium, K5-I SX, then K6-I onward for the compiled Hessian, compact output, and control-regularity choices. This makes the structural elimination visible early rather than presenting it as a late consequence of compilation.")
paragraph(doc, "The historical reduced isokinetic transcription keeps theta and omega as NLP states and enforces theta-dot = omega-reference and omega-dot = 0. The prescribed-kinematics alternative removes these equality-constrained variables and reconstructs them at every physical Radau or IRK stage from the same clock and phase origin.")
equation(doc, "θ(t) = θ₀ + ωref(t − t₀),     ω(t) = ωref")
paragraph(doc, "For four muscles, this changes the state dimension from 23 to 21: the 20 Ding states and Eprod remain. Fibre geometry, force-length, force-velocity, passive force, inverse load, and work rate are evaluated at the reconstructed exact stage angle. No lookup-table approximation or stage interpolation is introduced. The phase origin must be transferred consistently between RHO windows; resetting local time without the matching phase would change the Hill relations.")
paragraph(doc, "This is an elimination of prescribed kinematic equalities within one isokinetic transcription. It is not a comparison with the isoresistive task. Bounds, seed transfer, result summaries, and dense audits reconstruct theta and omega from the physical clock; the mode is restricted to reduced isokinetic mechanics and does not use redundant free-cadence velocity guards.")
paragraph(doc, "A new controlled pre-SX MX pair was run for 100 one-cycle windows. The current isokinetic driver uses periodic calcium, so it is a K4-I realization of the K3-I0/K3-I1 design: it establishes the kinematic elimination before SX, but does not isolate an interaction with the historical Ding representation. The cases were sequential, with 16 CasADi threads and single-thread numerical libraries, without explicit CPU pinning. They differ physically only in the kinematic representation.")
table(doc,
    ["Case", "Representation", "Certified", "Hot solve median / P90", "Complete RHO median / P90", "NLP variables"],
    [
        ["K3-I0 / K4-I reference", "MX, theta/omega states", "100/100", "4.631 / 5.487 s", "4.791 / 5.655 s", "4284"],
        ["K3-I1 / K4-I", "MX, prescribed", "100/100", "3.275 / 3.917 s", "3.435 / 4.106 s", "3922"],
    ], [1.0, 1.25, 0.65, 1.15, 1.3, 0.65])
paragraph(doc, "Eliminating the two prescribed kinematic variables removes 362 NLP variables in this transcription. It reduced the MX hot median by 29.3 percent and the complete-iteration median by 28.3 percent. Over the 100 saved windows, maximum states-versus-prescribed differences were 8.33e-6 microseconds in pulse width, 3.95e-10 rad in reconstructed angle, 7.68e-12 rad/s in speed, 6.22e-9 J in energy, and 3.09e-11 in relative capacity. Both campaigns passed 100/100 isokinetic and mechanical audits; the state-kinematics maximum effective primal infeasibility was 9.09e-7. The evidence is therefore an exact equality elimination for these trajectories, not a change to the muscle or mechanical model.")
table(doc,
    ["Measure", "State kinematics", "Prescribed kinematics", "Interpretation"],
    [
        ["NLP states, four muscles", "23", "21", "Ding states and Eprod retained"],
        ["IPOPT hot solve, median / P90", "0.707 / 0.808 s", "0.373 / 0.439 s", "100 windows, MA57 SX Radau-5"],
        ["Complete RHO iteration, median / P90", "0.797 / 0.899 s", "0.452 / 0.519 s", "Transfer, solver, certification"],
        ["IPOPT certification", "100 / 100", "100 / 100", "Isokinetic and mechanical audits passed"],
        ["Acados prescribed IRK 4x5", "--", "100 / 100; 0.202 / 0.207 s solve", "0.236 / 0.242 s iteration; four SQP iterations; IPOPT cycle-1 seed"],
    ], [1.15, 1.2, 1.75, 2.55])
paragraph(doc, "The SX production pair is the later K5-I0/K5-I1 continuation: in the matched 100-window IPOPT campaign, prescribed kinematics reduced hot solve time by 47.2 percent and complete RHO iteration time by 43.3 percent. Objective sums differed by 7.83e-8 and executed fatigue by 7.58e-8. The maximum saved-trace differences were 1.40e-11 microseconds in pulse width, 1.83e-8 J in energy, and 3.95e-10 rad in reconstructed angle. Historical K6–K9 records remain labelled state kinematics; they must be repeated from K5-I1 before being claimed as a cumulative prescribed sequence. Acados also certified 100/100 windows, but its internal objective and fatigue differed materially from IPOPT (986.147 / 953.999 versus 1552.602 / 1502.703). Its timing establishes robust feasibility and execution, not cross-solver optimality equivalence; a matched Acados state-kinematics run or common dense replay remains required.")

paragraph(doc, "7. References", "Heading 1")
for reference in [
    "[1] Ding J, Wexler AS, Binder-Macleod SA. A predictive model of fatigue in human skeletal muscles. Journal of Applied Physiology. 2000;89(4):1322-1332. doi:10.1152/jappl.2000.89.4.1322.",
    "[2] Ding J, Wexler AS, Binder-Macleod SA. A mathematical model that predicts the force-frequency relationship of human skeletal muscle. Muscle & Nerve. 2002;26(4):477-485. doi:10.1002/mus.10198.",
    "[3] Ding J, Wexler AS, Binder-Macleod SA. Mathematical models for fatigue minimization during functional electrical stimulation. Journal of Electromyography and Kinesiology. 2003;13(6):575-588. doi:10.1016/S1050-6411(03)00102-0.",
    "[4] Hairer E, Wanner G. Solving Ordinary Differential Equations II: Stiff and Differential-Algebraic Problems. 2nd ed. Springer; 1996. doi:10.1007/978-3-642-05221-7.",
    "[5] Hairer E, Wanner G. Stiff differential equations solved by Radau methods. Journal of Computational and Applied Mathematics. 1999;111(1-2):93-111. doi:10.1016/S0377-0427(99)00134-X.",
    "[6] De Groote F, Kinney AL, Rao AV, Fregly BJ. Evaluation of direct collocation optimal control problem formulations for solving the muscle redundancy problem. Annals of Biomedical Engineering. 2016;44(10):2922-2936. doi:10.1007/s10439-016-1591-9.",
    "[7] Co K, Puchaud P, Moissenet F, Begon M. Maximizing Task Endurance through Muscle Fatigue Minimization with Consideration of Muscle Fatigability and Task Contribution: an in-Silico FES Study of Handcycling. Manuscript. 2026.",
]:
    reference_paragraph = paragraph(doc, reference)
    reference_paragraph.paragraph_format.space_after = Pt(1)
    for run in reference_paragraph.runs:
        run.font.size = Pt(8.5)

doc.save(OUT)
